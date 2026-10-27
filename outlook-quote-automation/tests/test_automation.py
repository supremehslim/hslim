import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from quote_automation.core import Item, extract_rules, price_with_margin, totals, recipient, compose
from quote_automation.services import GooglePrices, matching, Graph, ServiceError
from quote_automation.runner import State, Runner, ROOT
from quote_automation.workbook import make_workbook

MESSAGE={'id':'mail1','internetMessageId':'<mail1@example.test>','subject':'견적문의 ABC-100', 'from':{'emailAddress':{'address':'customer@example.test'}},'body':{'contentType':'text','content':'ABC-100 3개'},'hasAttachments':False}
CFG={'margin_percent':10,'margin_mode':'markup','search_price_vat':'included'}

class Tests(unittest.TestCase):
    def test_margin_and_tax(self):
        self.assertEqual(price_with_margin(11000, CFG), 11000) # VAT 제외 매입 10000 +10%
        self.assertEqual(price_with_margin(11000,{**CFG,'margin_mode':'gross_margin','margin_percent':20}),12500)
        self.assertEqual(totals([Item('a','a',3,11000)]),{'amounts':[33000],'supply':33000,'vat':3300,'total':36300})
        self.assertEqual(totals([Item('a','a',1,15)])['vat'],2) # Excel ROUND: 1.5→2
        with self.assertRaises(ValueError):price_with_margin(11000,{**CFG,'margin_percent':None})
        with self.assertRaises(ValueError):price_with_margin(11000,{**CFG,'search_price_vat':'unknown'})
    def test_extraction_and_recipient(self):
        q=extract_rules('ABC-100 / 테스트 제품 / 3\nXYZ-200 2개')
        self.assertEqual([x.quantity for x in q['items']],[3,2])
        self.assertIsNone(extract_rules('예쁜 제품 좀 견적 부탁해요')['items'][0].quantity)
        m={**MESSAGE,'replyTo':[{'emailAddress':{'address':'reply@example.test'}}]}
        self.assertEqual(recipient(m)['emailAddress']['address'],'reply@example.test')
        with self.assertRaises(ValueError):recipient({**m,'replyTo':m['replyTo']*2})
        mail=compose(MESSAGE,{},[Item('a','a',None)],'서명')
        self.assertIn('확인필요',mail['subject']);self.assertEqual(mail['ccRecipients'],[])
    def test_exact_match_and_search(self):
        item=Item('ABC-100','ABC-100 1TB',3)
        self.assertTrue(matching(item,{'title':'브랜드 ABC100 1TB 신품'}))
        for title in ['ABC-1000 1TB','ABC-100 2TB','ABC-100 1TB 중고','ABC-100 1TB 카드할인']:
            self.assertFalse(matching(item,{'title':title}),title)
        results={'shopping_results':[{'title':'ABC100 1TB','price':'₩22,000','extracted_price':22000,'product_link':'https://seller.test/1'},
            {'title':'ABC-100 1TB','price':'₩11,000','extracted_price':11000,'product_link':'https://seller.test/2'},
            {'title':'ABC-1000 1TB','price':'₩1,000','extracted_price':1000,'product_link':'https://seller.test/3'}]}
        with patch.dict('os.environ',{'SERPAPI_API_KEY':'fixture-only'}),patch('quote_automation.services.request_json',return_value=results):
            out=GooglePrices(CFG).price(item);self.assertEqual(out.unit,11000);self.assertEqual(out.evidence['url'],'https://seller.test/2')
            out=GooglePrices({**CFG,'search_price_vat':'unknown'}).price(Item('ABC-100','ABC-100 1TB',1));self.assertIsNone(out.unit)
        with patch.dict('os.environ',{'SERPAPI_API_KEY':'fixture-only'}),patch('quote_automation.services.request_json',side_effect=ServiceError('HTTP 429')):
            out=GooglePrices(CFG).price(Item('ABC-100','ABC-100',1));self.assertIsNone(out.unit);self.assertIn('확인필요',out.note)
    def test_workbook(self):
        with tempfile.TemporaryDirectory() as t:
            file=Path(t)/'quote.xlsx'
            items=[Item('a','=HYPERLINK("bad")',3,11000),Item('b','b',2,20000),Item('c','c',1,15)]
            out=make_workbook(ROOT/'templates/견적서_검토용.xlsx',file,{},items,'customer@example.test')
            self.assertTrue(out['verified']);self.assertEqual(out['total'],80317)
            w=load_workbook(file,data_only=True);self.assertEqual(w.active['E18'].value,80317);self.assertEqual(w.active['B7'].value,'customer@example.test');self.assertTrue(w.active['B11'].value.startswith("'="));self.assertEqual(w.active['A13'].value,'c');self.assertEqual(w.active['D5'].value,'임현섭 부장');w.close()
            out=make_workbook(ROOT/'templates/견적서_검토용.xlsx',file,{},[Item('a','a',3,11000),Item('b','미확인',2)],'customer@example.test')
            self.assertIsNone(out['total']);w=load_workbook(file,data_only=True);self.assertEqual(w.active['E17'].value,'확인필요');w.close()
    def test_graph_no_send_and_origin(self):
        g=Graph('fixture-token')
        with self.assertRaises(ServiceError):g.call('POST','/me/sendMail')
        with self.assertRaises(ServiceError):g.call('GET','https://evil.test/data')
    def test_resume_idempotency(self):
        class FakeGraph:
            created=0; completed=0; fail=False
            def inbox(self,since):return [MESSAGE]
            def create_reply(self,id):
                self.created+=1
                if self.fail: raise ServiceError('uncertain response')
                return {'id':'draft1'}
            def complete_draft(self,id,mail,file):
                self.completed+=1;return {'id':id,'webLink':'https://outlook.office.com/'}
        class Prices:
            def price(self,item):item.unit=11000;return item
        with tempfile.TemporaryDirectory() as t:
            state=State(Path(t)/'state.db');g=FakeGraph();runner=Runner({'extractor':'rules'},g,Prices(),state,Path(t))
            runner.poll();runner.poll();self.assertEqual(g.created,1);self.assertEqual(g.completed,1);self.assertEqual(state.statuses()[0][1],'done');state.db.close()
        with tempfile.TemporaryDirectory() as t:
            state=State(Path(t)/'state.db');g=FakeGraph();g.fail=True;runner=Runner({'extractor':'rules'},g,Prices(),state,Path(t))
            runner.poll();runner.poll();self.assertEqual(g.created,1);self.assertEqual(state.statuses()[0][1],'ambiguous');state.db.close()

if __name__=='__main__':unittest.main()
