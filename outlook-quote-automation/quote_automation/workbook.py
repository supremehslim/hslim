from copy import copy
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import os
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.workbook.properties import CalcProperties
from .core import KST, safe_text, totals

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

def cache_formulas(path, values):
    """수식은 유지하고 Python 검산값을 OOXML 수식 캐시에 기록합니다."""
    with ZipFile(path) as z: contents = {n: z.read(n) for n in z.namelist()}
    root = ET.fromstring(contents['xl/worksheets/sheet1.xml'])
    for cell in root.findall('.//s:c', NS):
        if cell.attrib['r'] not in values: continue
        val = values[cell.attrib['r']]
        old = cell.find('s:v', NS)
        if old is not None: cell.remove(old)
        v = ET.SubElement(cell, '{'+NS['s']+'}v')
        if isinstance(val, str): cell.set('t', 'str')
        else: cell.attrib.pop('t', None)
        v.text = str(val)
    contents['xl/worksheets/sheet1.xml'] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    tmp = Path(str(path)+'.tmp')
    with ZipFile(tmp, 'w', ZIP_DEFLATED) as z:
        for n, data in contents.items(): z.writestr(n, data)
    os.replace(tmp, path)

def make_workbook(template, destination, quote, items, email):
    destination = Path(destination); destination.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(template)
    if len(wb.worksheets) != 1 or wb.active.title != '견적서': raise ValueError('지원하지 않는 양식입니다.')
    ws = wb.active
    if ws['A10'].value != '제품번호\nPart No': raise ValueError('견적서 양식의 품목 헤더가 변경되었습니다.')
    n = len(items); extra = n - 1
    if not 1 <= n <= 100: raise ValueError('품목 수 범위는 1~100개입니다.')
    merges = [str(r) for r in ws.merged_cells.ranges if r.min_row > 11]
    from openpyxl.worksheet.cell_range import CellRange
    for m in merges: ws.unmerge_cells(m)
    # footer row height도 함께 이동시킵니다.
    heights = {i: copy(v) for i, v in ws.row_dimensions.items() if i > 11}
    if extra: ws.insert_rows(12, extra)
    for m in merges:
        r = CellRange(m); r.shift(row_shift=extra); ws.merge_cells(str(r))
    for i, dim in sorted(heights.items(), reverse=True):
        dim.index = i + extra; ws.row_dimensions[i + extra] = dim
    for row in range(11, 11+n):
        ws.row_dimensions[row].height = max(ws.row_dimensions[11].height or 30, 35)
        for col in range(1, 6):
            ws.cell(row, col)._style = copy(ws.cell(11, col)._style)
            ws.cell(row, col).alignment = copy(ws.cell(11, col).alignment)
    ws['D4'] = datetime.now(KST).replace(tzinfo=None)
    ws['B5'] = safe_text(quote.get('customer') or '확인필요')
    ws['B6'] = safe_text(quote.get('contact') or '담당자 확인필요')
    ws['B7'] = email; ws['B8'] = safe_text(quote.get('phone') or '확인필요')
    # 원본 공급자 정보는 보존하고 납기·결제조건은 확약하지 않습니다.
    ws['D8'] = '발행일로부터 30일 (참고 견적)'
    values = {}; t = totals(items)
    for idx, item in enumerate(items):
        r = 11+idx
        ws.cell(r, 1, safe_text(item.part))
        ws.cell(r, 2, safe_text(item.description))
        ws.cell(r, 3, item.quantity if item.quantity is not None else '확인필요')
        ws.cell(r, 4, item.unit if item.unit is not None else '확인필요')
        ws.cell(r, 5, f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r}),C{r}>0,D{r}>=0),ROUND(C{r}*D{r},0),"확인필요")')
        for col in (4, 5): ws.cell(r, col).number_format = '#,##0'
        values[f'E{r}'] = t['amounts'][idx] if t['amounts'][idx] is not None else '확인필요'
        src = item.evidence
        comment = '\n'.join(x for x in [item.note, '공개 판매가 검색 기반 참고 견적', src.get('url'), src.get('searched_at')] if x)
        ws.cell(r, 4).comment = Comment(comment, '견적 자동화')
    sr = 14+extra; vr = sr+1; tr = sr+2; last = 10+n
    ws.cell(sr, 5, f'=IF(COUNT(E11:E{last})={n},SUM(E11:E{last}),"확인필요")')
    ws.cell(vr, 5, f'=IF(ISNUMBER(E{sr}),ROUND(E{sr}*0.1,0),"확인필요")')
    ws.cell(tr, 5, f'=IF(AND(ISNUMBER(E{sr}),ISNUMBER(E{vr})),SUM(E{sr}:E{vr}),"확인필요")')
    for r, key in [(sr, 'supply'), (vr, 'vat'), (tr, 'total')]:
        values[f'E{r}'] = t[key] if t[key] is not None else '확인필요'
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    ws.print_area = f'A1:E{25+extra}'
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    wb.save(destination); wb.close()
    cache_formulas(destination, values)
    saved = load_workbook(destination, data_only=False)
    cached = load_workbook(destination, data_only=True)
    try:
        for coord, value in values.items():
            if not str(saved.active[coord].value).startswith('=') or cached.active[coord].value != value:
                raise ValueError('저장 후 수식·캐시 검산 불일치: '+coord)
        if saved.active['D5'].value != '임현섭 부장': raise ValueError('공급자 정보가 변경되었습니다.')
    finally: saved.close(); cached.close()
    return {'verified': True, 'method': 'Python Decimal 계산 + 저장 파일 수식·캐시 재열람 검산', **t}
