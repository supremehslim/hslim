from __future__ import annotations
import argparse
import base64
import csv
import dataclasses
from datetime import datetime, timedelta, timezone
import hashlib
import io
import json
import logging
import os
from pathlib import Path
import sqlite3
import time
from zipfile import ZipFile
from .core import Item, compose, extract_rules, now, plain, recipient
from .services import Graph, GooglePrices, ServiceError, acquire_token, extract
from .workbook import make_workbook

ROOT = Path(__file__).resolve().parent.parent
LOG = logging.getLogger('quote')

def load_config():
    try:
        from dotenv import load_dotenv
        load_dotenv(ROOT/'.env')
    except ImportError: pass
    p = ROOT/'config.json'
    if not p.exists(): raise ServiceError('config.example.json을 config.json으로 복사하고 설정하세요.')
    cfg = json.loads(p.read_text(encoding='utf-8-sig'))
    if int(cfg.get('poll_seconds', 120)) < 30: raise ServiceError('poll_seconds는 30초 이상으로 설정하세요.')
    return cfg

class Lock:
    def __init__(self, file): self.file = file
    def __enter__(self):
        self.f = open(self.file, 'a+b'); self.f.seek(0); self.f.write(b'0'); self.f.flush(); self.f.seek(0)
        try:
            if os.name == 'nt':
                import msvcrt
                msvcrt.locking(self.f.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(self.f.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError:
            self.f.close(); raise ServiceError('다른 자동화 프로세스가 실행 중입니다.') from None
        return self
    def __exit__(self, *args): self.f.close()

class State:
    def __init__(self, path):
        self.db = sqlite3.connect(path)
        self.db.execute('CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY,value TEXT)')
        self.db.execute('CREATE TABLE IF NOT EXISTS jobs (key TEXT PRIMARY KEY,message TEXT,phase TEXT,draft_id TEXT,error TEXT,updated TEXT)')
        self.db.commit()
    def get(self, key):
        row = self.db.execute('SELECT value FROM meta WHERE key=?', (key,)).fetchone()
        return row[0] if row else None
    def set(self, key, val):
        self.db.execute('INSERT OR REPLACE INTO meta VALUES (?,?)', (key, val)); self.db.commit()
    def add(self, message):
        key = hashlib.sha256((message.get('internetMessageId') or message['id']).encode()).hexdigest()[:24]
        self.db.execute('INSERT OR IGNORE INTO jobs VALUES (?,?,?,NULL,NULL,?)', (key, json.dumps(message), 'preparing', now())); self.db.commit()
    def pending(self):
        return list(self.db.execute("SELECT key,message,phase,draft_id FROM jobs WHERE phase IN ('preparing','drafting')"))
    def update(self, key, phase, draft_id=None, error=None):
        self.db.execute('UPDATE jobs SET phase=?,draft_id=COALESCE(?,draft_id),error=?,updated=? WHERE key=?', (phase,draft_id,error,now(),key)); self.db.commit()
    def statuses(self): return list(self.db.execute('SELECT key,phase,draft_id,error,updated FROM jobs ORDER BY updated DESC LIMIT 50'))

def attachment_text(graph, message):
    if not message.get('hasAttachments'): return '', []
    from urllib.parse import quote
    p = '/me/messages/'+quote(message['id'], safe='')+'/attachments'
    parts = []; unresolved = []
    attachments = list(graph.pages(p, {'$select': 'id,name,size,isInline,contentType'}))
    for a in attachments:
        if a.get('isInline'): continue
        name = a.get('name', '이름없는 첨부')
        if a.get('size', 0) > 3_000_000:
            unresolved.append(name+': 3MB 초과'); continue
        suffix = Path(name).suffix.lower()
        if suffix not in ('.xlsx', '.csv', '.txt', '.pdf'):
            unresolved.append(name+': 미지원 형식'); continue
        try:
            detail = graph.call('GET', p+'/'+quote(a['id'], safe=''))
            b = base64.b64decode(detail.get('contentBytes', ''), validate=True)
            if not b or len(b) > 3_000_000: raise ValueError('attachment size')
            if suffix == '.xlsx':
                from openpyxl import load_workbook
                with ZipFile(io.BytesIO(b)) as z:
                    if sum(x.file_size for x in z.infolist()) > 20_000_000: raise ValueError('expanded size')
                w = load_workbook(io.BytesIO(b), read_only=True, data_only=True); rows = []
                try:
                    for s in w:
                        if s.max_row > 200 or s.max_column > 20: raise ValueError('row limit')
                        rows.extend(' | '.join(str(v) if v is not None else '' for v in row) for row in s.iter_rows(values_only=True))
                finally: w.close()
                text = '\n'.join(rows)
            elif suffix == '.pdf':
                from pypdf import PdfReader
                pdf = PdfReader(io.BytesIO(b))
                if len(pdf.pages) > 20: raise ValueError('page limit')
                text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
            else:
                try: text = b.decode('utf-8-sig')
                except UnicodeDecodeError: text = b.decode('cp949')
                if suffix == '.csv': text = '\n'.join(' | '.join(row) for row in csv.reader(io.StringIO(text)))
            if not text.strip() or len(text) > 24000: raise ValueError('empty or oversized text')
            parts.append('\n첨부자료 '+name+'\n'+text)
        except ServiceError:
            raise
        except Exception:
            unresolved.append(name+': 내용 추출 확인필요')
    if sum(map(len, parts)) > 24000: return '', unresolved+['첨부 전체 텍스트 길이 초과']
    return '\n'.join(parts), unresolved

class Runner:
    def __init__(self, cfg, graph, prices, state, directory=ROOT/'runtime'):
        self.cfg=cfg; self.graph=graph; self.prices=prices; self.state=state; self.directory=Path(directory)
    def poll(self):
        since = self.state.get('since')
        if not since:
            since = (datetime.now(timezone.utc)-timedelta(hours=float(self.cfg.get('initial_lookback_hours', 0)))).isoformat().replace('+00:00', 'Z')
            self.state.set('since', since)
        # 스캔 시작 시각을 커서로 보존하고 60초 중첩 조회해 수신 도중인 메일을 놓치지 않습니다.
        next_cursor = (datetime.now(timezone.utc)-timedelta(seconds=60)).isoformat().replace('+00:00', 'Z')
        for m in self.graph.inbox(since):
            if m.get('isDraft'): continue
            text = (m.get('subject', '')+'\n'+plain(m)).lower()
            if any(word.lower() in text for word in self.cfg.get('subject_keywords', ['견적', 'quotation', 'rfq'])):
                self.state.add(m)
        self.state.set('since', max(since, next_cursor))
        for key, message_json, phase, draft in self.state.pending():
            try: self.process(key, json.loads(message_json), phase, draft)
            except Exception as e:
                # 수신 메일/본문/인증값은 로그에 기록하지 않습니다.
                LOG.error('작업 %s 실패 (%s)', key, type(e).__name__)
                existing = self.state.db.execute('SELECT phase,draft_id FROM jobs WHERE key=?', (key,)).fetchone()
                if existing[0] == 'creating': self.state.update(key, 'ambiguous', error='초안 생성 응답 불명확. 임시보관함을 확인한 뒤 수동 조치 필요.')
                else: self.state.update(key, existing[0], existing[1], '작업 실패: '+type(e).__name__)
    def process(self, key, message, phase, draft):
        to = recipient(message)['emailAddress']['address']
        folder = self.directory/'quotes'/key; folder.mkdir(parents=True, exist_ok=True)
        attachment = folder/'견적서.xlsx'; mailfile = folder/'draft.json'
        if phase == 'preparing':
            text, missing = attachment_text(self.graph, message)
            try:
                quote = extract(message.get('subject', '')+'\n'+plain(message)+'\n'+text, self.cfg)
            except ServiceError:
                quote = extract_rules(''); quote['notes'] = '품목 추출 오류로 전체 요청 확인필요'
            items = quote.pop('items')
            if quote.get('notes'):
                items.append(Item('', '요청사항 확인필요', None, note=quote['notes']))
            if missing: items.append(Item('', '첨부 요청사항 확인필요', None, note='; '.join(missing)))
            if len(items) > 100: items = items[:99] + [Item('', '처리 한도 초과 품목 확인필요', None, note='100개 초과 항목은 수동 검토가 필요합니다.')]
            items = [self.prices.price(x) for x in items]
            verified = make_workbook(ROOT/'templates/견적서_검토용.xlsx', attachment, quote, items, to)
            mail = compose(message, quote, items, self.cfg.get('signature', ''))
            mailfile.write_text(json.dumps(mail, ensure_ascii=False, indent=2), encoding='utf-8')
            (folder/'audit.json').write_text(json.dumps({'generated_at': now(), 'message_id': message['id'], 'items': [dataclasses.asdict(x) for x in items], 'verification': verified}, ensure_ascii=False, indent=2), encoding='utf-8')
            self.state.update(key, 'creating')  # 불확실한 POST 재시도로 중복 초안이 생기지 않도록 먼저 기록
            reply = self.graph.create_reply(message['id'])
            draft = reply['id']; self.state.update(key, 'drafting', draft)
        mail = json.loads(mailfile.read_text(encoding='utf-8'))
        if attachment.stat().st_size > 2_500_000: raise ValueError('생성 견적서가 첨부 제한 크기를 초과했습니다.')
        result = self.graph.complete_draft(draft, mail, attachment)
        (folder/'saved-draft.json').write_text(json.dumps({'id': result['id'], 'webLink': result.get('webLink'), 'saved_at': now()},ensure_ascii=False),encoding='utf-8')
        self.state.update(key, 'done', draft)
        LOG.info('작업 %s: 견적서 첨부·수신자·제목·임시보관함 저장 검증 완료', key)

def demo():
    folder=ROOT/'runtime/demo'; folder.mkdir(parents=True,exist_ok=True)
    items=[Item('TEST-002', '자동화 동작 확인용 가상 품목 (실제 가격 아님)', 3, 11000, '데모', '테스트 데이터: 실제 검색가가 아닙니다.'), Item('UNKNOWN-001', '검색 결과 없는 경우의 표시 예시', 2)]
    q={'customer': '데모 고객사', 'contact': '테스트 담당자', 'phone': ''}
    m={'subject': '[데모] 견적문의', 'from': {'emailAddress': {'address': 'customer@example.test', 'name': '테스트 담당자'}}}
    result=make_workbook(ROOT/'templates/견적서_검토용.xlsx',folder/'데모_견적서.xlsx',q,items,'customer@example.test')
    mail=compose(m,q,items,'임현섭 부장\n주식회사 제이앤지시스템')
    (folder/'데모_메일.html').write_text('<!doctype html><meta charset="utf-8"><h1>데모 · 실제 메일 아님</h1><p>수신: customer@example.test</p><h2>'+mail['subject']+'</h2>'+mail['body']['content'],encoding='utf-8')
    print('데모 생성 완료:',folder); print('검산:',result['method'],result['verified'])

def main():
    parser=argparse.ArgumentParser(description='M365 Outlook 견적서·회신 초안 자동화 (자동 발송 없음)')
    parser.add_argument('command', choices=['auth','once','watch','demo','status','doctor','recover','reset'])
    parser.add_argument('--job', help='status에 표시되는 작업 ID')
    parser.add_argument('--draft', help='복구할 Outlook 초안 ID')
    parser.add_argument('--confirm-no-draft', action='store_true', help='Outlook에서 기존 초안이 없음을 직접 확인한 경우만 재생성 허용')
    args=parser.parse_args()
    if args.command=='demo': demo(); return
    cfg=load_config(); runtime=ROOT/'runtime'; runtime.mkdir(exist_ok=True)
    logging.basicConfig(level=logging.INFO,format='%(asctime)s %(levelname)s %(message)s',handlers=[logging.StreamHandler(),logging.FileHandler(runtime/'automation.log',encoding='utf-8')])
    if args.command=='doctor':
        for key in ['M365_CLIENT_ID','M365_TENANT_ID','SERPAPI_API_KEY','OPENAI_API_KEY']:
            print(key, '설정됨' if os.getenv(key) else '미설정')
        print('마진:',cfg.get('margin_percent'),'방식:',cfg.get('margin_mode'),'검색가 VAT:',cfg.get('search_price_vat'))
        print('기본 규칙 추출은 정형 메일만 지원합니다. 자유형은 OPENAI_API_KEY가 필요합니다.'); return
    with Lock(runtime/'process.lock'):
        state=State(runtime/'state.sqlite3')
        if args.command=='status':
            for row in state.statuses(): print(*row,sep=' | ')
            return
        if args.command in ('recover','reset'):
            row = state.db.execute('SELECT key,message,phase,draft_id FROM jobs WHERE key=?', (args.job,)).fetchone()
            if not row or row[2] not in ('creating','ambiguous','drafting'):
                raise ServiceError('복구 가능한 작업 ID가 아닙니다. status를 확인하세요.')
            if args.command=='reset':
                if not args.confirm_no_draft: raise ServiceError('기존 초안이 없음을 확인한 뒤 --confirm-no-draft를 지정하세요.')
                state.db.execute("UPDATE jobs SET phase='preparing',draft_id=NULL,error=NULL WHERE key=?", (args.job,));state.db.commit()
                print('재생성 대기 상태로 변경했습니다. 다음 once/watch 실행에서 처리합니다.');return
            if not args.draft: raise ServiceError('--draft로 기존 초안 ID를 지정하세요.')
            from urllib.parse import quote as urlquote
            graph=Graph(acquire_token(cfg,runtime))
            result=graph.call('GET','/me/messages/'+urlquote(args.draft,safe=''),query={'$select':'id,isDraft,toRecipients'})
            expected=json.loads((runtime/'quotes'/args.job/'draft.json').read_text(encoding='utf-8'))
            actual=[r['emailAddress']['address'].lower() for r in result.get('toRecipients',[])]
            target=[r['emailAddress']['address'].lower() for r in expected['toRecipients']]
            if not result.get('isDraft') or actual!=target: raise ServiceError('초안 상태 또는 수신자가 일치하지 않아 복구하지 않았습니다.')
            state.update(args.job,'drafting',args.draft)
            Runner(cfg,graph,GooglePrices(cfg),state,runtime).process(row[0],json.loads(row[1]),'drafting',args.draft)
            return
        if args.command=='auth': acquire_token(cfg,runtime,interactive=True); print('M365 로그인 완료. 자동 발송 권한은 요청하지 않았습니다.'); return
        while True:
            try:
                token=acquire_token(cfg,runtime)
                Runner(cfg,Graph(token),GooglePrices(cfg),state,runtime).poll()
            except ServiceError as e:
                LOG.error('%s',e)
                if args.command=='once': raise
            if args.command=='once': break
            time.sleep(int(cfg.get('poll_seconds',120)))

if __name__=='__main__': main()
